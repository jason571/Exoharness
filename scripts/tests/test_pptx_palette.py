"""Run in the existing Skill OCI image with /skill and writable /work mounts."""
import json
from pathlib import Path
import sys
import tempfile
import unittest

sys.path.insert(0, '/skill/scripts')
import render_deck
from check_layout import broken_text_code, inspect
from render_deck import ContentError, render


class PaletteTests(unittest.TestCase):
    def test_invalid_palette_diagnostic_identifies_exact_contract_and_retains_content(self):
        for palette, message in [({'primary': '142D40'}, 'unsupported key'),
                                 ({'ink': '#142D40'}, 'without #'),
                                 ({'accent': ['007F86']}, 'without #')]:
            with self.subTest(palette=palette), tempfile.TemporaryDirectory(dir='/work') as directory:
                output = Path(directory) / 'report.pptx'
                data = {'title': 'Exact saved input', 'palette': palette,
                        'slides': [{'layout': 'cover', 'title': '中国经济', 'subtitle': '已确认的统计期'}]}
                with self.assertRaisesRegex(ContentError, message):
                    render(data, output)
                self.assertEqual(json.loads(output.with_name('.report.pptx.content.json').read_text()), data)
                self.assertFalse(output.exists())

    def test_default_and_supported_palette_render_real_pptx(self):
        for palette in [None, {'ink': '142D40', 'accent': '007f86'}]:
            with self.subTest(palette=palette), tempfile.TemporaryDirectory(dir='/work') as directory:
                output = Path(directory) / 'report.pptx'
                data = {'slides': [{'layout': 'cover', 'title': '中国经济', 'subtitle': '已确认的统计期'}]}
                if palette is not None:
                    data['palette'] = palette
                result = render(data, output)
                self.assertTrue(output.is_file())
                self.assertEqual(result['slides'], 1)


class TextIntegrityTests(unittest.TestCase):
    def test_renderer_imports_the_shared_text_checker(self):
        self.assertIs(render_deck.broken_text_code, broken_text_code)

    def test_valid_unicode_and_whitespace_are_accepted_without_rewriting(self):
        for text in ('', '中国经济与人工智能', 'English 123', 'café Über Ångström',
                     'Ελληνικά العربية', '🤖 👨\u200d👩\u200d👧', '价格 £20 / €15',
                     'https://example.test/a?x=1&y=2', '第一行\n第二行\r\n\tAI'):
            with self.subTest(text=text):
                self.assertIsNone(broken_text_code(text))

    def test_replacement_controls_and_invalid_unicode_are_rejected(self):
        for text, code in [('损坏\ufffd文字', 'replacement_character'),
                           ('bad\x00text', 'control_character'),
                           ('bad\x07text', 'control_character'),
                           ('bad\x7ftext', 'control_character'),
                           ('bad\x85text', 'control_character'),
                           ('bad\ud800text', 'invalid_unicode_surrogate'),
                           ('bad\udffftext', 'invalid_unicode_surrogate')]:
            with self.subTest(text=repr(text)):
                self.assertEqual(broken_text_code(text), code)

    def test_reversible_cjk_mojibake_is_rejected_in_mixed_text(self):
        for encoding in ('latin-1', 'cp1252'):
            damaged = '中'.encode('utf-8').decode(encoding)
            for text in (damaged, 'prefix ' + damaged, '正常标题：' + damaged + '，尾注'):
                with self.subTest(encoding=encoding, text=text):
                    self.assertEqual(broken_text_code(text), 'reversible_cjk_mojibake')

    def test_damaged_copy_is_saved_exactly_without_publishing_a_deck(self):
        for text in ('损坏\ufffd文字', 'bad\x00text', '中'.encode('utf-8').decode('latin-1')):
            with self.subTest(text=repr(text)), tempfile.TemporaryDirectory(dir='/work') as directory:
                output = Path(directory) / 'report.pptx'
                data = {'slides': [{'layout': 'cover', 'title': text}]}
                with self.assertRaisesRegex(ContentError, 'damaged text encoding'):
                    render(data, output)
                self.assertEqual(json.loads(output.with_name('.report.pptx.content.json').read_text()), data)
                self.assertFalse(output.exists())

    def test_multilingual_copy_survives_actual_pptx_generation(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory(dir='/work') as directory:
            output = Path(directory) / 'report.pptx'
            data = {'slides': [{'layout': 'cover', 'title': '中文 AI', 'subtitle': 'café €20 · 🤖'}]}
            result = render(data, output)
            self.assertEqual(result['slides'], 1)
            deck = Presentation(output)
            text = '\n'.join(s.text for s in deck.slides[0].shapes if s.has_text_frame)
            self.assertIn('中文 AI', text)
            self.assertIn('café €20 · 🤖', text)
            self.assertEqual(json.loads(output.with_name('.report.pptx.content.json').read_text()), data)
            self.assertEqual(inspect(output)['status'], 'passed')

    def test_existing_geometry_rejection_is_preserved(self):
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory(dir='/work') as directory:
            output = Path(directory) / 'invalid.pptx'
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            shape = slide.shapes.add_textbox(Inches(-1), Inches(1), Inches(2), Inches(1))
            shape.text = '中文 text'
            deck.save(output)
            result = inspect(output)
            self.assertEqual(result['status'], 'failed')
            self.assertTrue(any(error['code'] == 'text_outside_slide' for error in result['errors']))


if __name__ == '__main__':
    unittest.main()
